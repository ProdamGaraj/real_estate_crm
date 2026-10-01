# -*- coding: utf-8 -*-
"""
Сборщик проектной документации: Markdown -> DOCX по ГОСТ.

Исходные тексты документов лежат в docs/gost/src/*.md, реквизиты (организации,
коды, утверждающие лица) — в docs/gost/meta.json. Сборщик формирует для
каждого документа файл .docx с титульным листом, листом утверждения
(для документов ЕСПД), содержанием, нумерацией разделов, таблиц и
приложений, колонтитулом с номером страницы и обозначением документа и
листом регистрации изменений.

    python docs/gost/tools/build.py            # все документы
    python docs/gost/tools/build.py 03-P3      # один документ (по имени файла)

Оглавление и число листов — поля Word. Их значения рассчитывает Word:
на Windows с установленным Word это делает tools/finalize.ps1 (он же
выгружает PDF), в остальных случаях — клавиша F9 после открытия файла.

Формат исходных текстов описан в docs/gost/README.md.
"""

import io
import json
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SRC = ROOT / 'src'
OUT = ROOT / 'out'
DATA = ROOT / 'data'
REPO = ROOT.parents[1]

# --- Геометрия листа: А4, поля для документов без рамки -----------------------
PAGE_W, PAGE_H = Mm(210), Mm(297)
MARGIN_LEFT, MARGIN_RIGHT = Mm(30), Mm(15)
MARGIN_TOP, MARGIN_BOTTOM = Mm(20), Mm(20)
TEXT_WIDTH_MM = 210 - 30 - 15

FONT = 'Times New Roman'
MONO = 'Courier New'
BODY_PT = 12
TABLE_PT = 10
INDENT = Cm(1.25)

# Буквы для перечислений и приложений по ГОСТ 2.105-2019:
# без ё, з, й, о, ч, ъ, ы, ь — их легко спутать с цифрами и другими буквами
LIST_LETTERS = 'абвгдежиклмнпрстуфхцшщэюя'
APPENDIX_LETTERS = 'АБВГДЕЖИКЛМНПРСТУФХЦШЩЭЮЯ'


# =============================================================================
# Разбор исходного текста
# =============================================================================

class Block:
    def __init__(self, kind, **kw):
        self.kind = kind
        self.__dict__.update(kw)


def parse_front_matter(text):
    meta = {}
    if text.startswith('---'):
        end = text.index('\n---', 3)
        for line in text[3:end].strip().splitlines():
            if ':' in line:
                k, v = line.split(':', 1)
                meta[k.strip()] = v.strip()
        text = text[end + 4:]
    return meta, text


LIST_RE = re.compile(r'^(\s*)(-|\d+\.)\s+(.*)$')


def parse_blocks(text):
    lines = text.splitlines()
    blocks = []
    i = 0
    n = len(lines)
    pending_caption = None
    pending_widths = None

    def is_block_start(s):
        return (s.startswith('#') or s.startswith('|') or s.startswith('```')
                or s.startswith('::: ') or s.startswith('> ') or s.strip() == '\\pagebreak'
                or s.startswith('Таблица:') or s.startswith('Ширины:') or s.startswith('Рисунок:')
                or bool(LIST_RE.match(s)))

    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped == '\\pagebreak':
            blocks.append(Block('pagebreak'))
            i += 1
            continue

        if line.startswith('#'):
            m = re.match(r'^(#+)\s+(.*)$', line)
            level = len(m.group(1))
            title = m.group(2).strip()
            if title.startswith('!'):
                blocks.append(Block('heading', level=level, title=title[1:].strip(), mode='struct'))
            elif title.startswith('@'):
                parts = [p.strip() for p in title[1:].split('|')]
                blocks.append(Block('heading', level=1, title=parts[2] if len(parts) > 2 else '',
                                    mode='appendix', status=parts[1] if len(parts) > 1 else 'обязательное'))
            else:
                blocks.append(Block('heading', level=level, title=title, mode='num'))
            i += 1
            continue

        if stripped.startswith('Рисунок:'):
            parts = [x.strip() for x in stripped.split(':', 1)[1].split('|')]
            blocks.append(Block('figure', caption=parts[0], path=parts[1],
                                width=float(parts[2]) if len(parts) > 2 else 150))
            i += 1
            continue

        if stripped.startswith('Таблица:'):
            pending_caption = stripped.split(':', 1)[1].strip()
            i += 1
            continue
        if stripped.startswith('Ширины:'):
            pending_widths = [float(x) for x in stripped.split(':', 1)[1].split(',')]
            i += 1
            continue

        if stripped.startswith('|'):
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                row = lines[i].strip()
                cells = [c.strip() for c in row.strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c):
                    rows.append(cells)
                i += 1
            blocks.append(Block('table', caption=pending_caption, widths=pending_widths,
                                header=rows[0], rows=rows[1:]))
            pending_caption = pending_widths = None
            continue

        if stripped.startswith('```'):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith('```'):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(Block('code', lines=code))
            continue

        if stripped.startswith('::: '):
            parts = stripped[4:].split()
            blocks.append(Block('include', name=parts[0], args=parts[1:]))
            i += 1
            continue

        if stripped.startswith('> '):
            note = []
            while i < n and lines[i].strip().startswith('> '):
                note.append(lines[i].strip()[2:])
                i += 1
            blocks.append(Block('note', text=' '.join(note)))
            continue

        m = LIST_RE.match(line)
        if m:
            items = []
            while i < n:
                m = LIST_RE.match(lines[i])
                if m:
                    level = 1 if len(m.group(1)) < 2 else 2
                    items.append({'level': level, 'ordered': m.group(2) != '-', 'text': m.group(3)})
                    i += 1
                    continue
                # Продолжение пункта — строка с отступом без маркера
                if lines[i].startswith('  ') and lines[i].strip() and items:
                    items[-1]['text'] += ' ' + lines[i].strip()
                    i += 1
                    continue
                break
            blocks.append(Block('list', items=items))
            continue

        para = [stripped]
        i += 1
        while i < n and lines[i].strip() and not is_block_start(lines[i]):
            para.append(lines[i].strip())
            i += 1
        blocks.append(Block('para', text=' '.join(para)))

    return blocks


INLINE_RE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*)')


def add_inline(paragraph, text, size=None, bold=None):
    """Добавляет текст с разметкой **жирный**, *курсив*, `моноширинный` и <br>."""
    for chunk_i, chunk in enumerate(text.split('<br>')):
        if chunk_i:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        for part in INLINE_RE.split(chunk):
            if not part:
                continue
            run_bold, run_italic, mono = bold, None, False
            if part.startswith('**') and part.endswith('**') and len(part) > 4:
                part, run_bold = part[2:-2], True
            elif part.startswith('`') and part.endswith('`') and len(part) > 2:
                part, mono = part[1:-1], True
            elif part.startswith('*') and part.endswith('*') and len(part) > 2:
                part, run_italic = part[1:-1], True
            run = paragraph.add_run(part)
            if run_bold:
                run.bold = True
            if run_italic:
                run.italic = True
            if mono:
                set_run_font(run, MONO)
                run.font.size = Pt((size or BODY_PT) - 1)
            elif size:
                run.font.size = Pt(size)


# =============================================================================
# Низкоуровневые помощники python-docx
# =============================================================================

def set_run_font(run, name):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(attr), name)


def style_font(style, name=FONT, size=BODY_PT, bold=False, italic=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(attr), name)
    for theme_attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if fonts.get(qn(theme_attr)) is not None:
            del fonts.attrib[qn(theme_attr)]
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rpr.append(lang)
    lang.set(qn('w:val'), 'ru-RU')
    lang.set(qn('w:eastAsia'), 'ru-RU')


def para_format(style_or_par, align=None, first=None, left=None, before=0, after=0, line=1.5,
                keep_next=None):
    pf = style_or_par.paragraph_format
    if align is not None:
        pf.alignment = align
    if first is not None:
        pf.first_line_indent = first
    if left is not None:
        pf.left_indent = left
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line == 1:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    else:
        pf.line_spacing = line
    if keep_next is not None:
        pf.keep_with_next = keep_next


def add_field(paragraph, instr, placeholder='1', dirty=False):
    """Сложное поле Word: begin / instrText / separate / результат / end."""
    def fld(kind):
        r = OxmlElement('w:r')
        f = OxmlElement('w:fldChar')
        f.set(qn('w:fldCharType'), kind)
        if kind == 'begin' and dirty:
            f.set(qn('w:dirty'), 'true')
        r.append(f)
        return r

    p = paragraph._p
    p.append(fld('begin'))
    r = OxmlElement('w:r')
    t = OxmlElement('w:instrText')
    t.set(qn('xml:space'), 'preserve')
    t.text = f' {instr} '
    r.append(t)
    p.append(r)
    p.append(fld('separate'))
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = placeholder
    r.append(t)
    p.append(r)
    p.append(fld('end'))


def set_cell_borders(cell, **edges):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcpr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        val = edges.get(edge, 'nil')
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), val)
        if val != 'nil':
            el.set(qn('w:sz'), '4')
            el.set(qn('w:color'), '000000')
        borders.append(el)


def no_borders(table):
    tbl = table._tbl
    tblpr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'nil')
        borders.append(el)
    tblpr.append(borders)


def set_col_widths(table, widths_mm):
    tbl = table._tbl
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths_mm):
        gc.set(qn('w:w'), str(int(Mm(w).twips)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_mm):
            cell.width = Mm(w)
    tblpr = tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblpr.append(layout)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trpr.append(el)


def cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    el.set(qn('w:val'), 'true')
    trpr.append(el)


def cell_margins(table, mm=1.5):
    tblpr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for edge in ('left', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:w'), str(int(Mm(mm).twips)))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblpr.append(mar)


# =============================================================================
# Стили документа
# =============================================================================

def setup_styles(doc):
    styles = doc.styles

    normal = styles['Normal']
    style_font(normal)
    para_format(normal, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=INDENT, line=1.5)
    normal.paragraph_format.widow_control = True

    # Заголовки разделов: с абзацного отступа, без точки в конце, полужирные
    sizes = {1: 14, 2: 12, 3: 12, 4: 12}
    for level in (1, 2, 3, 4):
        st = styles[f'Heading {level}']
        style_font(st, size=sizes[level], bold=True)
        para_format(st, align=WD_ALIGN_PARAGRAPH.LEFT, first=INDENT, left=Cm(0),
                    before=12 if level == 1 else 12, after=12 if level == 1 else 6,
                    line=1.5, keep_next=True)
        st.paragraph_format.widow_control = True
        # Без переноса слов в заголовках
        ppr = st.element.get_or_add_pPr()
        sup = OxmlElement('w:suppressAutoHyphens')
        ppr.append(sup)

    def new_style(name, base='Normal', **kw):
        try:
            st = styles[name]
        except KeyError:
            st = styles.add_style(name, 1)
            st.base_style = styles[base]
        return st

    t = new_style('GostTable')
    style_font(t, size=TABLE_PT)
    para_format(t, align=WD_ALIGN_PARAGRAPH.LEFT, first=Cm(0), line=1)

    th = new_style('GostTableHead')
    style_font(th, size=TABLE_PT, bold=True)
    para_format(th, align=WD_ALIGN_PARAGRAPH.CENTER, first=Cm(0), line=1)

    cap = new_style('GostCaption')
    style_font(cap)
    para_format(cap, align=WD_ALIGN_PARAGRAPH.LEFT, first=Cm(0), before=6, after=0, line=1.5,
                keep_next=True)

    code = new_style('GostCode')
    style_font(code, name=MONO, size=9)
    para_format(code, align=WD_ALIGN_PARAGRAPH.LEFT, first=Cm(0), left=INDENT, line=1)

    center = new_style('GostCenter')
    style_font(center)
    para_format(center, align=WD_ALIGN_PARAGRAPH.CENTER, first=Cm(0), line=1.5)

    struct = new_style('GostStruct')
    style_font(struct, size=14, bold=True)
    para_format(struct, align=WD_ALIGN_PARAGRAPH.CENTER, first=Cm(0), before=0, after=12,
                line=1.5, keep_next=True)

    title = new_style('GostTitle')
    style_font(title)
    para_format(title, align=WD_ALIGN_PARAGRAPH.LEFT, first=Cm(0), line=1)

    note = new_style('GostNote')
    style_font(note, size=11)
    para_format(note, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=INDENT, line=1.15, before=3, after=3)

    for name in ('GostCenter', 'GostStruct', 'GostTitle', 'GostCaption', 'GostTable', 'GostTableHead'):
        ppr = styles[name].element.get_or_add_pPr()
        ppr.append(OxmlElement('w:suppressAutoHyphens'))

    # Оглавление: без абзацного отступа, точки до номера страницы
    for level in (1, 2, 3):
        st = new_style(f'TOC {level}')
        style_font(st)
        para_format(st, align=WD_ALIGN_PARAGRAPH.LEFT, first=Cm(0),
                    left=Cm(0.5 * (level - 1)), line=1.15)
        st.paragraph_format.tab_stops.add_tab_stop(
            Mm(TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    # Язык документа — русский: проверка орфографии и переносы
    settings = doc.settings.element
    hyph = OxmlElement('w:autoHyphenation')
    hyph.set(qn('w:val'), 'true')
    settings.append(hyph)


def setup_section(section):
    section.page_width, section.page_height = PAGE_W, PAGE_H
    section.left_margin, section.right_margin = MARGIN_LEFT, MARGIN_RIGHT
    section.top_margin, section.bottom_margin = MARGIN_TOP, MARGIN_BOTTOM
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)


def restart_numbering(section, start=1):
    sectpr = section._sectPr
    pg = sectpr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType')
        sectpr.append(pg)
    pg.set(qn('w:start'), str(start))


def clear_story(story):
    for p in list(story.paragraphs):
        p._p.getparent().remove(p._p)


# =============================================================================
# Документ
# =============================================================================

class GostDocument:
    def __init__(self, meta, info, includes):
        self.meta = meta          # реквизиты комплекта
        self.info = info          # шапка конкретного документа
        self.includes = includes
        self.doc = Document()
        setup_styles(self.doc)
        setup_section(self.doc.sections[0])
        self.counters = [0, 0, 0, 0]
        self.table_no = 0
        self.figure_no = 0
        self.appendix_idx = -1    # -1 — основная часть
        self.first_heading = True
        self.kind = info.get('kind', 'gost34')
        self.designation = designation(meta, info)

    # --- служебное -----------------------------------------------------------
    def p(self, text='', style=None, align=None):
        par = self.doc.add_paragraph(style=style)
        if text:
            add_inline(par, text)
        if align is not None:
            par.alignment = align
        return par

    def page_break(self):
        par = self.doc.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.add_run().add_break(WD_BREAK.PAGE)

    # --- колонтитул: номер страницы и обозначение ---------------------------
    def setup_header(self, section):
        section.different_first_page_header_footer = True
        header = section.header
        header.is_linked_to_previous = False
        clear_story(header)
        par = header.add_paragraph(style='GostCenter')
        par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        add_field(par, 'PAGE', '2')
        par2 = header.add_paragraph(style='GostCenter')
        par2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        add_inline(par2, self.designation, size=11)
        first = section.first_page_header
        first.is_linked_to_previous = False
        clear_story(first)
        first.add_paragraph()
        footer = section.footer
        footer.is_linked_to_previous = False
        first_footer = section.first_page_footer
        first_footer.is_linked_to_previous = False

    # --- титульная часть -----------------------------------------------------
    def approval_table(self, left_title, left_lines, right_title, right_lines):
        table = self.doc.add_table(rows=1, cols=2)
        no_borders(table)
        set_col_widths(table, [TEXT_WIDTH_MM / 2, TEXT_WIDTH_MM / 2])
        for cell, title, body in ((table.cell(0, 0), left_title, left_lines),
                                  (table.cell(0, 1), right_title, right_lines)):
            cell.paragraphs[0].style = self.doc.styles['GostTitle']
            if title:
                r = cell.paragraphs[0].add_run(title)
                r.bold = True
            for line in body:
                par = cell.add_paragraph(style='GostTitle')
                par.paragraph_format.space_before = Pt(4)
                add_inline(par, line)
        return table

    def signature_block(self, position, org, name):
        year = self.meta['year']
        return [position, org, f'_____________ {name}', f'«___» ____________ {year} г.']

    def title_center(self, lines_before=6):
        for _ in range(lines_before):
            self.p(style='GostCenter')

    def title_bottom(self):
        par = self.p(f'{self.meta["city"]} {self.meta["year"]}', style='GostCenter')
        return par

    def write_title_gost34(self):
        m, info = self.meta, self.info
        cust, dev = m['customer'], m['developer']
        self.approval_table(
            'СОГЛАСОВАНО', self.signature_block(dev['approver_position'], dev['name'], dev['approver_name']),
            'УТВЕРЖДАЮ', self.signature_block(cust['approver_position'], cust['name'], cust['approver_name']),
        )
        self.title_center(5)
        par = self.p(style='GostCenter')
        r = par.add_run(m['system']['full_name'].upper())
        r.bold = True
        par = self.p(f'«{m["system"]["short_name"]}»', style='GostCenter')
        par.runs[0].bold = True
        self.p(style='GostCenter')
        par = self.p(style='GostCenter')
        r = par.add_run(info['title'].upper())
        r.bold = True
        r.font.size = Pt(16)
        if info.get('subtitle'):
            self.p(info['subtitle'], style='GostCenter')
        self.p(style='GostCenter')
        self.p(self.designation, style='GostCenter')
        par = self.p(style='GostCenter')
        par.add_run('Листов ')
        add_field(par, 'NUMPAGES', '__')
        if info.get('effective'):
            self.p(f'Действует с {info["effective"]}', style='GostCenter')
        self.title_center(6)
        self.title_bottom()

    def write_lu(self):
        """Лист утверждения документа ЕСПД (ГОСТ 19.104-78)."""
        m, info = self.meta, self.info
        cust, dev = m['customer'], m['developer']
        self.approval_table(
            '', [],
            'УТВЕРЖДАЮ', self.signature_block(cust['approver_position'], cust['name'], cust['approver_name']),
        )
        self.title_center(4)
        par = self.p(style='GostCenter')
        r = par.add_run(m['system']['full_name'].upper())
        r.bold = True
        par = self.p(f'«{m["system"]["short_name"]}»', style='GostCenter')
        par.runs[0].bold = True
        self.p(style='GostCenter')
        par = self.p(style='GostCenter')
        r = par.add_run(info['title'].upper())
        r.bold = True
        r.font.size = Pt(16)
        self.p(style='GostCenter')
        par = self.p(style='GostCenter')
        r = par.add_run('ЛИСТ УТВЕРЖДЕНИЯ')
        r.bold = True
        self.p(f'{self.designation}-ЛУ', style='GostCenter')
        self.title_center(3)
        self.approval_table(
            'СОГЛАСОВАНО', self.signature_block(dev['approver_position'], dev['name'], dev['approver_name']),
            'РАЗРАБОТЧИКИ', self.signature_block(dev['author_position'], dev['name'], dev['author_name']),
        )
        self.title_center(3)
        self.title_bottom()

    def write_title_espd(self):
        m, info = self.meta, self.info
        par = self.p(style='GostTitle')
        par.add_run('УТВЕРЖДЕН').bold = True
        self.p(f'{self.designation}-ЛУ', style='GostTitle')
        self.title_center(9)
        par = self.p(style='GostCenter')
        r = par.add_run(m['system']['full_name'].upper())
        r.bold = True
        par = self.p(f'«{m["system"]["short_name"]}»', style='GostCenter')
        par.runs[0].bold = True
        self.p(style='GostCenter')
        par = self.p(style='GostCenter')
        r = par.add_run(info['title'].upper())
        r.bold = True
        r.font.size = Pt(16)
        self.p(style='GostCenter')
        self.p(self.designation, style='GostCenter')
        par = self.p(style='GostCenter')
        par.add_run('Листов ')
        add_field(par, 'SECTIONPAGES', '__')
        self.title_center(10)
        self.title_bottom()

    def write_front(self):
        """Титульная часть, аннотация и содержание."""
        doc = self.doc
        if self.kind == 'espd':
            # Лист утверждения — отдельный документ: свой раздел без нумерации
            self.write_lu()
            doc.add_section(WD_SECTION.NEW_PAGE)
            # После add_section объекты разделов берём заново: python-docx
            # переносит свойства прежнего раздела в новый элемент
            lu, body = doc.sections[0], doc.sections[1]
            setup_section(body)
            restart_numbering(body, 1)
            self.setup_header(body)
            # У листа утверждения колонтитулов нет
            for story in (lu.header, lu.first_page_header, lu.footer):
                story.is_linked_to_previous = False
                clear_story(story)
                story.add_paragraph()
            self.write_title_espd()
        else:
            self.setup_header(doc.sections[0])
            restart_numbering(doc.sections[0], 1)
            self.write_title_gost34()

        if self.info.get('annotation'):
            self.page_break()
            self.p('АННОТАЦИЯ', style='GostStruct')
            for chunk in self.info['annotation'].split(' || '):
                self.p(chunk)

        self.page_break()
        self.p('СОДЕРЖАНИЕ', style='GostStruct')
        par = self.p(style='Normal')
        par.paragraph_format.first_line_indent = Cm(0)
        add_field(par, 'TOC \\o "1-3" \\h \\z \\u',
                  'Содержание формируется полем Word: выделите его и нажмите F9.')

    # --- основная часть ------------------------------------------------------
    def heading(self, b):
        doc = self.doc
        if b.mode == 'appendix':
            self.appendix_idx += 1
            self.counters = [0, 0, 0, 0]
            self.table_no = 0
            self.figure_no = 0
            letter = APPENDIX_LETTERS[self.appendix_idx]
            par = doc.add_paragraph(style='Heading 1')
            par.paragraph_format.page_break_before = True
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.first_line_indent = Cm(0)
            par.add_run(f'Приложение {letter}')
            par.add_run().add_break(WD_BREAK.LINE)
            r = par.add_run(f'({b.status})')
            r.bold = False
            par.add_run().add_break(WD_BREAK.LINE)
            par.add_run(b.title)
            return

        if b.mode == 'struct':
            par = doc.add_paragraph(style='Heading 1')
            par.paragraph_format.page_break_before = True
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.first_line_indent = Cm(0)
            par.add_run(b.title.upper())
            return

        level = b.level
        if self.appendix_idx >= 0:
            # В приложении заголовок ## — первый уровень нумерации (А.1)
            level = max(1, level - 1)
        self.counters[level - 1] += 1
        for k in range(level, 4):
            self.counters[k] = 0
        nums = [str(c) for c in self.counters[:level]]
        prefix = '.'.join(nums)
        if self.appendix_idx >= 0:
            prefix = APPENDIX_LETTERS[self.appendix_idx] + '.' + prefix
        style_level = min(level + (1 if self.appendix_idx >= 0 else 0), 4)
        par = doc.add_paragraph(style=f'Heading {style_level}')
        if level == 1 and self.appendix_idx < 0:
            par.paragraph_format.page_break_before = True
        add_inline(par, f'{prefix} {b.title}')

    def para(self, b):
        self.p(b.text)

    def note(self, b):
        par = self.doc.add_paragraph(style='GostNote')
        add_inline(par, b.text)

    def code(self, b):
        for i, line in enumerate(b.lines or ['']):
            par = self.doc.add_paragraph(style='GostCode')
            par.add_run(line.replace('\t', '    ') or ' ')
            if i == 0:
                par.paragraph_format.space_before = Pt(3)
        if b.lines:
            par.paragraph_format.space_after = Pt(6)

    def list(self, b):
        counters = {1: 0, 2: 0}
        for item in b.items:
            level = item['level']
            if level == 1:
                counters[2] = 0
            counters[level] += 1
            if level == 1:
                label = (LIST_LETTERS[counters[1] - 1] + ')') if item['ordered'] else '–'
            else:
                label = f'{counters[2]})'
            par = self.doc.add_paragraph(style='Normal')
            par.paragraph_format.first_line_indent = INDENT + Cm(0.75 * (level - 1))
            add_inline(par, f'{label} {item["text"]}')

    def table(self, b, widths=None, font=None):
        header, rows = b.header, b.rows
        cols = len(header)
        if b.caption:
            self.table_no += 1
            num = str(self.table_no)
            if self.appendix_idx >= 0:
                num = APPENDIX_LETTERS[self.appendix_idx] + '.' + num
            cap = self.doc.add_paragraph(style='GostCaption')
            add_inline(cap, f'Таблица {num} – {b.caption}')
        table = self.doc.add_table(rows=1 + len(rows), cols=cols)
        table.style = self.doc.styles['Table Grid']
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_margins(table)
        weights = widths or b.widths or [1] * cols
        total = sum(weights)
        widths_mm = [TEXT_WIDTH_MM * w / total for w in weights]
        set_col_widths(table, widths_mm)
        for ci, text in enumerate(header):
            cell = table.cell(0, ci)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            par = cell.paragraphs[0]
            par.style = self.doc.styles['GostTableHead']
            add_inline(par, text)
        repeat_header(table.rows[0])
        cant_split(table.rows[0])
        for ri, row in enumerate(rows, start=1):
            cant_split(table.rows[ri])
            for ci in range(cols):
                text = row[ci] if ci < len(row) else ''
                cell = table.cell(ri, ci)
                par = cell.paragraphs[0]
                par.style = self.doc.styles['GostTable']
                add_inline(par, text, size=font or TABLE_PT)
        # Отступ после таблицы
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        return table

    def figure(self, b):
        self.figure_no += 1
        num = str(self.figure_no)
        if self.appendix_idx >= 0:
            num = APPENDIX_LETTERS[self.appendix_idx] + '.' + num
        par = self.doc.add_paragraph(style='GostCenter')
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.space_before = Pt(6)
        par.add_run().add_picture(str(ROOT / b.path), width=Mm(b.width))
        cap = self.doc.add_paragraph(style='GostCenter')
        cap.paragraph_format.space_after = Pt(6)
        add_inline(cap, f'Рисунок {num} – {b.caption}')

    # --- лист регистрации изменений -----------------------------------------
    def change_sheet(self):
        par = self.doc.add_paragraph(style='Heading 1')
        par.paragraph_format.page_break_before = True
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Cm(0)
        par.add_run('ЛИСТ РЕГИСТРАЦИИ ИЗМЕНЕНИЙ')
        cols = 10
        table = self.doc.add_table(rows=2 + 22, cols=cols)
        table.style = self.doc.styles['Table Grid']
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_margins(table, 0.8)
        set_col_widths(table, [10, 16, 16, 16, 18, 20, 16, 23, 15, 15])
        heads = ['Изм.', 'Номера листов (страниц)', '', '', '', 'Всего листов (страниц) в докум.',
                 '№ докум.', 'Входящий № сопроводи-тельного докум. и дата', 'Подп.', 'Дата']
        subs = ['изменённых', 'заменённых', 'новых', 'аннулиро-ванных']
        for ci, text in enumerate(heads):
            par = table.cell(0, ci).paragraphs[0]
            par.style = self.doc.styles['GostTableHead']
            par.add_run(text).font.size = Pt(9)
        for k, text in enumerate(subs):
            par = table.cell(1, 1 + k).paragraphs[0]
            par.style = self.doc.styles['GostTableHead']
            par.add_run(text).font.size = Pt(9)
        table.cell(0, 1).merge(table.cell(0, 4))
        for ci in (0, 5, 6, 7, 8, 9):
            table.cell(0, ci).merge(table.cell(1, ci))
        for ri in range(2, 24):
            table.rows[ri].height = Mm(8)
        repeat_header(table.rows[0])
        repeat_header(table.rows[1])

    # --- сборка --------------------------------------------------------------
    def build(self, blocks):
        self.write_front()
        for b in blocks:
            if b.kind == 'heading':
                self.heading(b)
            elif b.kind == 'para':
                self.para(b)
            elif b.kind == 'list':
                self.list(b)
            elif b.kind == 'table':
                self.table(b)
            elif b.kind == 'code':
                self.code(b)
            elif b.kind == 'note':
                self.note(b)
            elif b.kind == 'pagebreak':
                self.page_break()
            elif b.kind == 'figure':
                self.figure(b)
            elif b.kind == 'include':
                fn = self.includes.get(b.name)
                if fn is None:
                    raise SystemExit(f'Неизвестная вставка ::: {b.name}')
                for sub in fn(self, *b.args):
                    {'heading': self.heading, 'para': self.para, 'list': self.list,
                     'table': self.table, 'code': self.code, 'note': self.note,
                     'pagebreak': lambda _b: self.page_break()}[sub.kind](sub)
        self.change_sheet()
        props = self.doc.core_properties
        props.title = f'{self.meta["system"]["short_name"]}. {self.info["title"]}'
        props.subject = self.designation
        props.author = self.meta['developer']['name']
        props.language = 'ru-RU'
        return self.doc


def designation(meta, info):
    c = meta['codes']
    if info.get('kind') == 'espd':
        return (f'{c["espd_country"]}.{c["espd_org"]}.{c["espd_reg"]}-{c["edition"]} '
                f'{info["code"]} {info.get("number", "01")}')
    return f'{c["gost34_org"]}.{c["gost34_class"]}.{c["gost34_reg"]}.{info["code"]}'


# =============================================================================
# Вставки, формируемые из данных проекта
# =============================================================================

def load_json(name):
    return json.load(io.open(DATA / name, encoding='utf-8'))


SQL_TYPES = {
    'BigAutoField': 'bigint, ПК',
    'AutoField': 'integer, ПК',
    'CharField': 'varchar({max_length})',
    'SlugField': 'varchar({max_length})',
    'EmailField': 'varchar({max_length})',
    'URLField': 'varchar({max_length})',
    'TextField': 'text',
    'IntegerField': 'integer',
    'BigIntegerField': 'bigint',
    'SmallIntegerField': 'smallint',
    'PositiveIntegerField': 'integer ≥ 0',
    'PositiveSmallIntegerField': 'smallint ≥ 0',
    'PositiveBigIntegerField': 'bigint ≥ 0',
    'BooleanField': 'boolean',
    'DateField': 'date',
    'DateTimeField': 'timestamptz',
    'TimeField': 'time',
    'DurationField': 'interval',
    'DecimalField': 'numeric({max_digits},{decimal_places})',
    'FloatField': 'double precision',
    'JSONField': 'jsonb',
    'FileField': 'varchar({max_length}), путь к файлу',
    'ImageField': 'varchar({max_length}), путь к файлу',
    'GenericIPAddressField': 'inet',
    'UUIDField': 'uuid',
}


def sql_type(f):
    t = f['type']
    if t in ('ForeignKey', 'OneToOneField'):
        suffix = ', уникальн.' if t == 'OneToOneField' else ''
        return f'bigint, ВК → {f["related_table"]}{suffix}'
    if t == 'ManyToManyField':
        return f'связь М:М через {f["through_table"]}'
    tmpl = SQL_TYPES.get(t, t)
    return tmpl.format(**{k: f.get(k) for k in ('max_length', 'max_digits', 'decimal_places')})


def describe_field(f):
    text = f['verbose'][:1].upper() + f['verbose'][1:] if f['verbose'] else ''
    extras = []
    if f.get('choices') and f['type'] != 'BooleanField':
        extras.append('значения: ' + ', '.join(code for code, _ in f['choices']))
    if f.get('on_delete') and f['type'] in ('ForeignKey', 'OneToOneField'):
        extras.append({'CASCADE': 'удаляется вместе со связанной записью',
                       'PROTECT': 'связанную запись удалить нельзя',
                       'SET_NULL': 'при удалении связанной записи обнуляется',
                       'SET_DEFAULT': 'при удалении связанной записи — значение по умолчанию',
                       'DO_NOTHING': 'целостность обеспечивает приложение'}.get(f['on_delete'], f['on_delete']))
    if f.get('unique') and not f.get('primary_key') and f['type'] != 'OneToOneField':
        extras.append('уникально')
    if extras:
        text = (text + '. ' if text else '') + '; '.join(extras)
    return text


def inc_catalog(gd, *args):
    """Каталог таблиц БД: таблица на каждую модель приложения."""
    models = load_json('models.json')
    app_names = {
        'realty': 'Каталог недвижимости', 'crm': 'Работа с клиентами', 'deals': 'Сделки',
        'finances': 'Финансы', 'tasks': 'Задачи', 'documents': 'Документы',
        'reports': 'Планы и отчёты', 'permissions': 'Организационная структура и права доступа',
    }
    out = []
    current_app = None
    for m in models:
        if m['app'] != current_app:
            current_app = m['app']
            out.append(Block('heading', level=2, title=f'{app_names.get(current_app, current_app)} '
                                                       f'(приложение {current_app})', mode='num'))
        out.append(Block('heading', level=3, title=f'Таблица {m["table"]} — {m["verbose"]}', mode='num'))
        notes = []
        if m['unique_together']:
            notes.append('Уникальные сочетания: ' + '; '.join('(' + ', '.join(u) + ')' for u in m['unique_together']) + '.')
        if m['constraints']:
            notes.append('Ограничения целостности: ' + ', '.join(m['constraints']) + '.')
        if m['indexes']:
            notes.append('Дополнительные индексы: ' + '; '.join(
                '(' + ', '.join(i['fields']) + ')' for i in m['indexes']) + '.')
        rows = []
        for f in m['fields']:
            column = f['column'] or f['name']
            required = 'нет' if (f['null'] or f['blank'] or f['type'] == 'ManyToManyField') else 'да'
            if f['primary_key']:
                required = 'да'
            rows.append([f'`{column}`', sql_type(f), required, describe_field(f)])
        out.append(Block('table', caption=f'Структура таблицы {m["table"]}',
                         widths=[26, 30, 11, 50], header=['Поле', 'Тип', 'Обяз.', 'Назначение'], rows=rows))
        for n in notes:
            out.append(Block('para', text=n))
    return out


def inc_tables_list(gd, *args):
    models = load_json('models.json')
    rows = [[f'`{m["table"]}`', m['verbose'][:1].upper() + m['verbose'][1:], m['app'], str(len(m['fields']))]
            for m in models]
    return [Block('table', caption='Перечень таблиц базы данных', widths=[48, 46, 16, 10],
                  header=['Таблица', 'Содержание', 'Приложение', 'Полей'], rows=rows)]


def inc_classifiers(gd, *args):
    """Классификаторы — перечисления, заданные в коде (TextChoices)."""
    models = load_json('models.json')
    out, seen = [], set()
    rows = []
    for m in models:
        for f in m['fields']:
            if not f.get('choices') or f['type'] == 'BooleanField':
                continue
            key = tuple(map(tuple, f['choices']))
            name = f'{m["table"]}.{f["column"] or f["name"]}'
            if key in seen and len(f['choices']) > 3:
                continue
            seen.add(key)
            values = '<br>'.join(f'`{code}` — {label}' for code, label in f['choices'])
            rows.append([f'`{name}`', f['verbose'][:1].upper() + f['verbose'][1:], values])
    out.append(Block('table', caption='Классификаторы, заданные в программе', widths=[38, 30, 52],
                     header=['Поле', 'Классификатор', 'Код — значение'], rows=rows))
    return out


def inc_routes(gd, *args):
    routes = load_json('routes.json')
    groups = [
        ('Служебные и аутентификация', ('/api/health', '/api/token', '/api/permissions/auth', '/api/permissions/me',
                                        '/api/schema', '/media')),
        ('Публичный API партнёров', ('/api/public',)),
        ('Клиенты, заявки, встречи', ('/api/clients', '/api/applications', '/api/application-statuses',
                                       '/api/meetings', '/api/rejection-reasons', '/api/users',
                                       '/api/dashboard')),
        ('Каталог недвижимости и скидки', ('/api/projects', '/api/buildings', '/api/building-types', '/api/discounts')),
        ('Сделки, платежи, документы', ('/api/deals', '/api/finances', '/api/templates')),
        ('Отчёты и планы', ('/api/reports',)),
        ('Задачи', ('/api/tasks', '/api/task-comments')),
        ('Организационная структура и права', ('/api/permissions',)),
    ]
    used = set()
    out = []
    for title, prefixes in groups:
        rows = []
        for r in routes:
            if r['path'] in ('/api/', '/api/permissions/'):
                continue
            if any(r['path'].startswith(p) for p in prefixes) and r['path'] not in used:
                used.add(r['path'])
                rows.append([f'`{r["path"]}`', ', '.join(r['methods']), r['doc'].rstrip('.') if r['doc'] and not r['doc'].startswith(('ViewSet', 'The default')) else ''])
        if rows:
            out.append(Block('table', caption=f'Маршруты API: {title.lower()}', widths=[62, 20, 38],
                             header=['Маршрут', 'Методы', 'Назначение'], rows=rows))
    return out


def git(*args):
    return subprocess.run(['git', *args], cwd=REPO, capture_output=True, text=True,
                          encoding='utf-8').stdout


def inc_manifest(gd, *args):
    """Опись исходных текстов: файлы под версионным контролем с хешами SHA-256."""
    import hashlib
    files = [f for f in git('ls-files').splitlines() if f]
    keep_ext = ('.py', '.ts', '.tsx', '.js', '.cjs', '.json', '.html', '.css', '.sh', '.yml', '.yaml',
                '.conf', '.txt', '.md', '.toml', '.cfg', '.ini', '.svg')
    skip_prefix = ('docs/', 'frontend-new/public/', 'backend/media/')
    skip_names = ('package-lock.json',)
    rows = []
    groups = {}
    for f in sorted(files):
        if f.startswith(skip_prefix) or Path(f).name in skip_names:
            continue
        if not f.endswith(keep_ext) and Path(f).name not in ('Dockerfile', '.env.example'):
            continue
        path = REPO / f
        if not path.exists():
            continue
        data = path.read_bytes()
        lines = data.count(b'\n') + (1 if data and not data.endswith(b'\n') else 0)
        sha = hashlib.sha256(data).hexdigest()
        top = f.split('/')[0] if '/' in f else '(корень)'
        groups.setdefault(top, [0, 0, 0])
        groups[top][0] += 1
        groups[top][1] += lines
        groups[top][2] += len(data)
        rows.append([f'`{f}`', str(lines), sha[:16]])
    summary = [[k, str(v[0]), str(v[1]), f'{v[2] / 1024:.0f}'] for k, v in sorted(groups.items())]
    total = [sum(v[i] for v in groups.values()) for i in range(3)]
    summary.append(['**Итого**', f'**{total[0]}**', f'**{total[1]}**', f'**{total[2] / 1024:.0f}**'])
    head = git('rev-parse', 'HEAD').strip()
    dirty = bool(git('status', '--porcelain', '--untracked-files=no').strip())
    gd.manifest_info = {'commit': head, 'dirty': dirty, 'files': total[0]}
    return [
        Block('table', caption='Сводные сведения об исходных текстах', widths=[40, 25, 30, 25],
              header=['Каталог', 'Файлов', 'Строк', 'Объём, КиБ'], rows=summary),
        Block('para', text=f'Опись составлена по коммиту `{head[:12]}` репозитория'
                           + (' с учётом незафиксированных изменений рабочего каталога.' if dirty else '.')),
        Block('table', caption='Опись файлов исходных текстов', widths=[78, 12, 30],
              header=['Файл', 'Строк', 'SHA-256 (первые 16 знаков)'], rows=rows),
    ]


def inc_documents(gd, *args):
    """Ведомость: перечень документов комплекта с обозначениями."""
    rows = []
    for path in sorted(SRC.glob('*.md')):
        info, _ = parse_front_matter(io.open(path, encoding='utf-8').read())
        if info.get('in_register', 'yes') == 'no':
            continue
        rows.append([designation(gd.meta, info), info['title'], info.get('standard', ''), '1', info.get('register_note', '')])
    return [Block('table', caption='Перечень документов', widths=[42, 40, 22, 9, 22],
                  header=['Обозначение', 'Наименование', 'Стандарт', 'Кол. экз.', 'Примечание'], rows=rows)]


def inc_snippet(gd, name, *args):
    """Общий фрагмент из src/_snippets/<name>.md — например, перечень терминов."""
    path = SRC / '_snippets' / f'{name}.md'
    _, body = parse_front_matter(io.open(path, encoding='utf-8').read())
    return parse_blocks(body)


INCLUDES = {
    'snippet': inc_snippet,
    'catalog': inc_catalog,
    'tables': inc_tables_list,
    'classifiers': inc_classifiers,
    'routes': inc_routes,
    'manifest': inc_manifest,
    'documents': inc_documents,
}


# =============================================================================

def build_one(path, meta):
    text = io.open(path, encoding='utf-8').read()
    info, body = parse_front_matter(text)
    blocks = parse_blocks(body)
    gd = GostDocument(meta, info, INCLUDES)
    doc = gd.build(blocks)
    OUT.mkdir(parents=True, exist_ok=True)
    out = OUT / f'{info["file"]}.docx'
    doc.save(out)
    note = ''
    if getattr(gd, 'manifest_info', None) and gd.manifest_info['dirty']:
        note = '  (внимание: в рабочем каталоге есть незафиксированные изменения — пересоберите после коммита)'
    print(f'  {gd.designation:34s} {out.name}{note}')
    return out


def main(argv):
    meta = json.load(io.open(ROOT / 'meta.json', encoding='utf-8'))
    if meta.get('year') == 'auto':
        meta['year'] = str(date.today().year)
    sources = sorted(SRC.glob('*.md'))
    if argv:
        sources = [s for s in sources if any(a in s.stem for a in argv)]
    print('Сборка документов:')
    for path in sources:
        build_one(path, meta)
    print('Готово. Оглавление и число листов рассчитает Word: tools/finalize.ps1 или F9.')


if __name__ == '__main__':
    main(sys.argv[1:])
