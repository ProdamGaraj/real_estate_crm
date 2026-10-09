"""
Формирование договора по шаблону.

Метка {{client.phone_number}} описана в справке по меткам и в инструкциях,
но телефон клиента давно хранится отдельной таблицей — метка подставляла
пустоту. Тест закрепляет, что в договор попадает основной номер клиента.
"""

import io
import shutil
import tempfile
from datetime import timedelta
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document
from rest_framework.test import APIClient

from apps.crm.models import Client, ClientPhoneNumber
from apps.deals.models import Deal
from apps.realty.models import Building, Project, Property
from permissions.models import Company, UserProfile

from .models import Template

User = get_user_model()
MEDIA_ROOT = tempfile.mkdtemp(prefix='crm-test-media-')


def docx_bytes(text):
    doc = Document()
    doc.add_paragraph(text)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def docx_text(content):
    return '\n'.join(p.text for p in Document(io.BytesIO(content)).paragraphs)


@override_settings(
    MEDIA_ROOT=MEDIA_ROOT,
    CACHES={'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache'}},
)
class GenerateDocumentBase(TestCase):
    """Сделка и шаблон для проверок формирования договора."""

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA_ROOT, ignore_errors=True)

    def setUp(self):
        self.company = Company.objects.create(name='Компания', code='DOCS')
        self.user = User.objects.create_superuser('admin', 'admin@example.com', 'pass-for-tests')
        profile, _ = UserProfile.objects.get_or_create(user=self.user)
        profile.company = self.company
        profile.is_system_admin = True
        profile.save()
        self.api = APIClient()
        self.api.force_authenticate(self.user)

        project = Project.objects.create(name='ЖК Тест', address='Ташкент', company=self.company)
        building = Building.objects.create(project=project, name='1', floors_count=9)
        prop = Property.objects.create(
            building=building, property_type='APARTMENT', unit_number='12', floor=3,
            area=Decimal('50.00'), price=Decimal('500000000'),
        )
        self.client_obj = Client.objects.create(full_name='Иванов Иван', company=self.company)
        self.deal = Deal.objects.create(
            company=self.company, client=self.client_obj, property=prop,
            booking_end_date=timezone.now() + timedelta(days=7),
            initial_price=Decimal('500000000'), initial_price_per_sqm=Decimal('10000000'),
            contract_number='Д-1', contract_date=timezone.now().date(), created_by=self.user,
        )
        self.template = Template.objects.create(
            name='Договор', company=self.company,
            file=SimpleUploadedFile('contract.docx', docx_bytes('Телефон: {{client.phone_number}}')),
        )
        self.url = f'/api/deals/{self.deal.id}/generate-document/{self.template.id}/'

    def generate(self):
        response = self.api.get(self.url)
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        return docx_text(response.content)


class GenerateDocumentPhoneTests(GenerateDocumentBase):
    def test_primary_phone_is_substituted(self):
        ClientPhoneNumber.objects.create(client=self.client_obj, phone_number='90 111 22 33')
        ClientPhoneNumber.objects.create(client=self.client_obj, phone_number='+998 93 444 55 66', is_primary=True)
        self.assertIn('Телефон: +998934445566', self.generate())

    def test_first_phone_when_none_is_primary(self):
        ClientPhoneNumber.objects.create(client=self.client_obj, phone_number='90 111 22 33')
        ClientPhoneNumber.objects.create(client=self.client_obj, phone_number='93 444 55 66')
        self.assertIn('Телефон: +998901112233', self.generate())

    def test_client_without_phone_gets_empty_value(self):
        self.assertIn('Телефон: ', self.generate())


class AmountWordsTests(TestCase):
    """Сумма прописью, разряды и даты для меток договора."""

    def test_russian_numbers_agree_in_gender_and_case(self):
        from .formatting import number_words_ru
        self.assertEqual(number_words_ru(0), 'ноль')
        self.assertEqual(number_words_ru(21000), 'двадцать одна тысяча')
        self.assertEqual(number_words_ru(22000), 'двадцать две тысячи')
        self.assertEqual(number_words_ru(111000), 'сто одиннадцать тысяч')
        self.assertEqual(number_words_ru(2500001), 'два миллиона пятьсот тысяч один')
        self.assertEqual(
            number_words_ru(247231950),
            'двести сорок семь миллионов двести тридцать одна тысяча девятьсот пятьдесят',
        )

    def test_currency_forms(self):
        from .formatting import amount_words
        self.assertEqual(amount_words('140583150.54'),
                         'сто сорок миллионов пятьсот восемьдесят три тысячи сто пятьдесят сумов 54 тийина')
        self.assertEqual(amount_words(1, 'USD'), 'один доллар США')
        self.assertEqual(amount_words('21000.01', 'USD'), 'двадцать одна тысяча долларов США 01 цент')
        self.assertEqual(amount_words(None), '')

    def test_uzbek(self):
        from .formatting import amount_words_uz
        self.assertEqual(amount_words_uz('1250000'), "bir million ikki yuz ellik ming so'm")
        self.assertEqual(amount_words_uz(12840, 'USD'), "o'n ikki ming sakkiz yuz qirq AQSH dollari")

    def test_money_and_dates(self):
        from datetime import date
        from .formatting import NBSP, date_short, date_text, money
        self.assertEqual(money(Decimal('247231950.00')), f'247{NBSP}231{NBSP}950')
        self.assertEqual(money('140583150.54'), f'140{NBSP}583{NBSP}150,54')
        from .formatting import number
        self.assertEqual(number(Decimal('48.00')), '48')
        self.assertEqual(number(Decimal('120.50')), '120,5')
        self.assertEqual(date_short(date(2026, 10, 9)), '09.10.2026')
        self.assertEqual(date_text(date(2026, 10, 9)), '9 октября 2026 г.')


class GenerateDocumentFiltersTests(GenerateDocumentBase):
    """Фильтры и график в шаблоне; опечатка в метке — понятная ошибка, а не 500."""

    def use_template(self, text):
        self.template.file = SimpleUploadedFile('contract2.docx', docx_bytes(text))
        self.template.save()

    def test_price_in_words_and_dates(self):
        self.deal.contract_price = Decimal('247231950.00')
        self.deal.contract_date = timezone.datetime(2026, 10, 9).date()
        self.deal.save()
        self.use_template('{{ deal.contract_price|money }} ({{ deal.contract_price|amount_words }}), '
                          '{{ deal.contract_date|date }}, {{ deal.contract_date|date_text }}')
        text = self.generate().replace(' ', ' ')
        self.assertIn('247 231 950 (двести сорок семь миллионов двести тридцать одна тысяча '
                      'девятьсот пятьдесят сумов), 09.10.2026, 9 октября 2026 г.', text)

    def test_payment_schedule_loop(self):
        from apps.finances.models import BeneficiaryAccount, Payment, PaymentType
        ptype = PaymentType.objects.create(name='Рассрочка', company=self.company)
        account = BeneficiaryAccount.objects.create(name='Счёт', details='-', company=self.company)
        for days, amount in ((30, '1000000'), (0, '2000000')):
            Payment.objects.create(deal=self.deal, client=self.client_obj, company=self.company,
                                   amount=Decimal(amount), currency='UZS', payment_type=ptype,
                                   beneficiary_account=account,
                                   due_date=timezone.localdate() + timedelta(days=days))
        self.use_template('{% for p in payments %}{{ p.amount|money }};{% endfor %}')
        # По сроку: сначала платёж «сегодня», потом через месяц
        self.assertIn('2 000 000;1 000 000;', self.generate())

    def test_empty_field_is_blank_not_none(self):
        self.deal.contract_number = None
        self.deal.save()
        self.use_template('Договор №[{{ deal.contract_number }}], выдан [{{ client.passport_issued_date|date }}]')
        self.assertIn('Договор №[], выдан []', self.generate())

    def test_broken_tag_returns_readable_error(self):
        self.use_template('{{ deal.contract_price|no_such_filter }}')
        response = self.api.get(self.url)
        self.assertEqual(response.status_code, 400)
        self.assertIn('Ошибка в метках шаблона', response.data['error'])

    def test_payment_schedule_table(self):
        # Как в шпаргалке: строка с {%tr for %}, строка платежа, строка {%tr endfor %}
        from apps.finances.models import BeneficiaryAccount, Payment, PaymentType
        ptype = PaymentType.objects.create(name='Первый взнос', company=self.company)
        account = BeneficiaryAccount.objects.create(name='Счёт', details='-', company=self.company)
        for days, amount in ((0, '3000000'), (30, '1500000.50')):
            Payment.objects.create(deal=self.deal, client=self.client_obj, company=self.company,
                                   amount=Decimal(amount), currency='UZS', payment_type=ptype,
                                   beneficiary_account=account,
                                   due_date=timezone.localdate() + timedelta(days=days))
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).text = '{%tr for p in payments %}'
        table.cell(1, 0).text = '{{ loop.index }}'
        table.cell(1, 1).text = '{{ p.amount|money }}'
        table.cell(1, 2).text = '{{ p.payment_type.name }}'
        table.cell(2, 0).text = '{%tr endfor %}'
        stream = io.BytesIO()
        doc.save(stream)
        self.template.file = SimpleUploadedFile('table.docx', stream.getvalue())
        self.template.save()

        response = self.api.get(self.url)
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        rows = [[cell.text for cell in row.cells] for row in Document(io.BytesIO(response.content)).tables[0].rows]
        self.assertEqual(rows, [
            ['1', '3 000 000', 'Первый взнос'],
            ['2', '1 500 000,50', 'Первый взнос'],
        ])

    def test_uzbek_cyrillic_phones_and_first_payment(self):
        # Шаблон закалата: сумма прописью кириллицей, оба телефона, закалат — первый платёж
        from apps.finances.models import BeneficiaryAccount, Payment, PaymentType
        ClientPhoneNumber.objects.create(client=self.client_obj, phone_number='93 971 88 11')
        ClientPhoneNumber.objects.create(client=self.client_obj, phone_number='90 111 22 33', is_primary=True)
        ptype = PaymentType.objects.create(name='Закалат', company=self.company)
        account = BeneficiaryAccount.objects.create(name='Счёт', details='-', company=self.company)
        for days, amount in ((10, '94000000'), (0, '6000000')):
            Payment.objects.create(deal=self.deal, client=self.client_obj, company=self.company,
                                   amount=Decimal(amount), currency='UZS', payment_type=ptype,
                                   beneficiary_account=account,
                                   due_date=timezone.localdate() + timedelta(days=days))
        self.deal.contract_price = Decimal('100000000')
        self.deal.save()
        self.use_template('{{ deal.contract_price|money }} ({{ deal.contract_price|words_uz_cyr }}) сўм; '
                          '{{ first_payment.amount|money }} ({{ first_payment.amount|words_uz_cyr }}) сўм; '
                          'Тел: {{ client.all_phones }}')
        text = self.generate().replace(' ', ' ')
        self.assertIn('100 000 000 (юз миллион) сўм; 6 000 000 (олти миллион) сўм; '
                      'Тел: +998901112233 / +998939718811', text)
